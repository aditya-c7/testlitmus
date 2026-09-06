import csv
import io
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from pypdf import PdfReader

TEXT_SUFFIXES = {".txt", ".md"}


@dataclass(frozen=True)
class Document:
    citation: str
    category: str
    text: str


def load_corpus(corpus_dir: Path) -> list[Document]:
    documents: list[Document] = []
    for path in sorted(p for p in corpus_dir.rglob("*") if p.is_file()):
        text = _read_document(path)
        if not text.strip():
            continue
        citation = path.relative_to(corpus_dir).as_posix()
        documents.append(Document(citation=citation, category=path.parent.name, text=text))
    if not documents:
        raise RuntimeError(f"no readable documents found in {corpus_dir}")
    return documents


def _read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        return "\n".join(" | ".join(row) for row in rows)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx_bytes(path.read_bytes(), str(path))
    if suffix in {".xlsx", ".xls"}:
        return _read_spreadsheet(path)
    return ""


def extract_upload(filename: str, data: bytes) -> str:
    """Extract draft text from raw uploaded bytes (pdf, docx, or plain text)."""
    name = filename or ""
    suffix = Path(name).suffix.lower()
    head = bytes(data[:4])
    if suffix == ".pdf" or head.startswith(b"%PDF"):
        return _read_pdf_bytes(data)
    if suffix == ".docx" or head.startswith(b"PK"):
        return _read_docx_bytes(data, name)
    try:
        return bytes(data).decode("utf-8", errors="replace")
    except Exception:
        return ""


def _read_docx_bytes(data: bytes, name: str = "") -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed; run pip install -r requirements.txt"
        ) from exc
    document = docx.Document(io.BytesIO(bytes(data)))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(p for p in (parts or []) if p and p.strip())


def _read_pdf(path: Path) -> str:
    return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)


def _read_pdf_bytes(data: bytes) -> str:
    return "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(bytes(data))).pages)


def _read_spreadsheet(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = [
            " | ".join("" if cell is None else str(cell) for cell in row)
            for row in sheet.iter_rows(values_only=True)
        ]
        sheets.append(f"[sheet: {sheet.title}]\n" + "\n".join(rows))
    workbook.close()
    return "\n\n".join(sheets)
